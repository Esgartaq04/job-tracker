"""Cover letters (Phase 5b): the cleaned job description plus a resume uploaded for
this one call, turned into a letter that only claims what the resume shows.

The resume is never stored. The letter is, as a `documents` row of kind
`cover_letter`, one per listing, so it can be edited and downloaded later.
"""

import io
import re
import uuid

import sqlalchemy as sa
from sqlalchemy.orm import Session

from src.core.config import settings
from src.models import Application, Document
from src.models.util import utcnow
from src.services import ai

KIND = "cover_letter"

SYSTEM = (
    "You write cover letters for job applications. Use only facts from the candidate's "
    "resume: never invent employers, titles, dates, degrees, skills, numbers, or "
    "achievements, and never claim experience the resume doesn't show. Connect the "
    "candidate's real experience to what the job description asks for, specifically. "
    "Write 250-400 words in 3-4 paragraphs, professional and warm, without stock "
    "openers like 'I am writing to express my interest'. Address it to the hiring team "
    "at the company and sign off with the candidate's name exactly as the resume gives "
    "it. Never leave placeholders or brackets — if a detail such as the hiring "
    "manager's name isn't known, write around it. Reply with only the letter: no "
    "subject line, no date or address block, no commentary before or after."
)


class NotReady(ValueError):
    """There's no description to write against yet."""


def letter_for(db: Session, application: Application) -> Document | None:
    return db.scalar(
        sa.select(Document)
        .where(Document.application_id == application.id, Document.kind == KIND)
        .order_by(Document.created_at.desc())
        .limit(1)
    )


def job_description(application: Application) -> str:
    """What the letter is written against: the user's own edit, else the cleaned text.
    The raw scrape is used only once the clean-up has run and declined, so a letter is
    never written against page chrome that a clean-up was about to remove."""
    if application.description_user:
        return application.description_user
    if application.description_clean:
        return application.description_clean
    cleanup = (application.extraction_meta or {}).get("cleanup") or {}
    if application.description_raw and cleanup.get("status") in {"rejected", "skipped"}:
        return application.description_raw
    if application.description_raw:
        raise NotReady("The description is still being cleaned up. Try again in a moment.")
    raise NotReady("This listing has no description yet. Add one in the Description tab.")


def generate(
    db: Session,
    application: Application,
    resume: bytes,
    *,
    filename: str | None = None,
    notes: str | None = None,
) -> Document:
    description = job_description(application)
    resume_blocks = ai.resume_content(resume, filename)
    client = ai.client()
    ai.reserve_call(db, application.user_id, KIND)

    job = (
        f"<job>\nCompany: {application.company or 'unknown'}\n"
        f"Role: {application.title or 'unknown'}\n"
        f"Location: {application.location or 'unknown'}\n\n{description}\n</job>"
    )
    ask = "Write my cover letter for this job."
    if notes and notes.strip():
        ask += f"\n\nThings I want it to mention:\n<notes>\n{notes.strip()[:2000]}\n</notes>"

    response = client.messages.create(
        model=settings.llm_model,
        max_tokens=4000,
        system=SYSTEM,
        messages=[
            {
                "role": "user",
                "content": [*resume_blocks, {"type": "text", "text": f"{job}\n\n{ask}"}],
            }
        ],
    )
    content = ai.response_text(response)

    letter = letter_for(db, application)
    if letter is None:
        letter = Document(
            user_id=application.user_id,
            application_id=application.id,
            kind=KIND,
            label=f"Cover letter — {application.company or 'company'}",
        )
        db.add(letter)
    letter.content = content
    letter.updated_at = utcnow()
    db.flush()
    return letter


def docx_bytes(content: str) -> bytes:
    import docx
    from docx.shared import Pt

    document = docx.Document()
    style = document.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)
    for paragraph in re.split(r"\n\s*\n", content.strip()):
        # Keep the line breaks inside a paragraph — a sign-off is "Sincerely,\nName".
        document.add_paragraph(paragraph.strip())
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def download_name(application: Application) -> str:
    parts = ["Cover Letter", application.company, application.title]
    name = " - ".join(part for part in parts if part)
    return re.sub(r"[^\w .,()&-]+", "", name).strip()[:120] + ".docx"


def owned_letter(db: Session, user_id: uuid.UUID, application: Application) -> Document | None:
    letter = letter_for(db, application)
    return letter if letter is not None and letter.user_id == user_id else None
