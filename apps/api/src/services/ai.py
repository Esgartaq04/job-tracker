"""Everything the AI features share: a client on the user's own Claude key, the usage
record, and how a resume gets into a request.

Four features call a model — Tier 4 extraction, description clean-up, cover letters,
and the extension's autofill — and each runs on the key of the user it's for. There is
no server-wide key. Each degrades the same ways: no key saved, a key Claude refused, or
the user's Claude account out of credit or rate-limited. Keeping the client here means
a new feature can't reach Claude on anyone else's key.
"""

import base64
import io
import logging
import uuid
import zipfile
from typing import Literal

from sqlalchemy.orm import Session

from src.models.llm_call import LLMCall
from src.services import ai_keys

logger = logging.getLogger(__name__)

#: Well past any real resume, well short of what a free instance wants in memory.
MAX_RESUME_BYTES = 5 * 1024 * 1024

# What the user is told. Each says what to do next, since the extension shows it as is.
MISSING_KEY = "Add your Claude API key (👤 → Claude API key) to use this."
REJECTED_KEY = "Claude rejected your saved key. Replace it under 👤 → Claude API key."
OUT_OF_CREDIT = (
    "Your Claude account is out of credit. Add credit at console.anthropic.com, then try again."
)
RATE_LIMITED = "Your Claude account hit its rate limit. Try again in a minute."
AI_CALL_FAILED = "The AI service couldn't complete that. Try again shortly."


class AIUnavailable(RuntimeError):
    """No usable key for this user, or the `anthropic` package isn't installed."""


class AIKeyMissing(AIUnavailable):
    """The user hasn't saved a key, or the saved one can't be opened."""


class AIKeyRejected(AIUnavailable):
    """Claude refused the user's saved key, now or on an earlier call."""


class AIFailed(RuntimeError):
    """The model answered, but not with something usable (a refusal, an empty reply)."""


class ResumeRejected(ValueError):
    def __init__(self, status_code: int, message: str) -> None:
        super().__init__(message)
        self.status_code = status_code


def user_key(db: Session, user_id: uuid.UUID) -> str:
    """The user's saved key, or why there isn't a usable one. A key Claude refused is
    refused here too, without asking Claude again, until the user saves a new one."""
    key = ai_keys.read_key(db, user_id)
    if key is None:
        raise AIKeyMissing(MISSING_KEY)
    if ai_keys.get_row(db, user_id).status == "rejected":
        raise AIKeyRejected(REJECTED_KEY)
    return key


def has_key(db: Session, user_id: uuid.UUID) -> bool:
    """For call sites that skip AI rather than fail without it."""
    try:
        user_key(db, user_id)
    except AIUnavailable:
        return False
    return True


def _build_client(api_key: str):
    try:
        import anthropic
    except ImportError as exc:
        raise AIUnavailable("AI features are off: the anthropic package isn't installed.") from exc
    return anthropic.Anthropic(api_key=api_key)


class UserClient:
    """An Anthropic client on one user's key. When Claude refuses the key mid-call, the
    key is marked rejected on the spot, so the next call stops at `user_key`."""

    def __init__(self, inner, db: Session, user_id: uuid.UUID) -> None:
        self._inner = inner
        self._db = db
        self._user_id = user_id
        self.messages = _Messages(self)

    def _call(self, method: str, request: dict):
        try:
            return getattr(self._inner.messages, method)(**request)
        except Exception as exc:
            if sdk_error_kind(exc) != "auth":
                raise
            logger.info("Claude rejected the saved key of user %s", self._user_id)
            ai_keys.mark_rejected(self._db, self._user_id)
            raise AIKeyRejected(REJECTED_KEY) from exc


class _Messages:
    """`client.messages`, as the call sites use it."""

    def __init__(self, client: UserClient) -> None:
        self._client = client

    def create(self, **request):
        return self._client._call("create", request)

    def parse(self, **request):
        return self._client._call("parse", request)


def client_for(db: Session, user_id: uuid.UUID) -> UserClient:
    """A client on the user's own key, built for this call. Nothing is cached, so a key
    never outlives the request or job that loaded it."""
    return UserClient(_build_client(user_key(db, user_id)), db, user_id)


def client_or_none(db: Session, user_id: uuid.UUID) -> UserClient | None:
    """`client_for`, or None when AI can't run for this user (Tier 4 just skips)."""
    try:
        return client_for(db, user_id)
    except AIUnavailable:
        return None


def sdk_error_kind(exc: BaseException) -> Literal["auth", "rate_limit", "credit", "other"] | None:
    """What an Anthropic SDK error says about the user's key or account, or None when
    `exc` isn't one. "auth" is their key; "rate_limit" and "credit" are their account's
    limits, not ours."""
    try:
        import anthropic
    except ImportError:
        return None
    if isinstance(exc, (anthropic.AuthenticationError, anthropic.PermissionDeniedError)):
        return "auth"
    if isinstance(exc, anthropic.RateLimitError):
        return "rate_limit"
    if isinstance(exc, anthropic.BadRequestError) and "credit balance" in str(exc).lower():
        return "credit"
    if isinstance(exc, anthropic.APIError):
        return "other"
    return None


def is_sdk_error(exc: Exception) -> bool:
    """An error from the API itself: rate limits, overload, timeouts, bad requests."""
    return sdk_error_kind(exc) is not None


def reserve_call(db: Session, user_id: uuid.UUID, kind: str) -> None:
    """Record a call before making it. Only a usage record: the limits are the user's
    own Claude account's."""
    db.add(LLMCall(user_id=user_id, kind=kind))
    db.flush()


def record_call(db: Session, user_id: uuid.UUID, kind: str) -> None:
    """Count a call that already happened (Tier 4 runs inside the pure pipeline)."""
    db.add(LLMCall(user_id=user_id, kind=kind))


def resume_content(data: bytes, filename: str | None = None) -> list[dict]:
    """Content blocks that put a resume in front of the model.

    A PDF goes in as a document block: the model reads its layout directly, so no PDF
    parser is needed here. A DOCX is unpacked to text. Nothing is written anywhere —
    the bytes live for the length of one request.
    """
    if not data:
        raise ResumeRejected(422, "The resume file is empty.")
    if len(data) > MAX_RESUME_BYTES:
        raise ResumeRejected(413, "That resume is over 5 MB. Export a smaller PDF and try again.")

    if data.startswith(b"%PDF"):
        return [
            {
                "type": "document",
                "source": {
                    "type": "base64",
                    "media_type": "application/pdf",
                    "data": base64.standard_b64encode(data).decode("ascii"),
                },
                "title": filename or "resume.pdf",
            }
        ]

    if _is_docx(data):
        text = docx_text(data)
        if not text.strip():
            raise ResumeRejected(422, "That Word file has no text in it.")
        return [{"type": "text", "text": f"<resume>\n{text}\n</resume>"}]

    raise ResumeRejected(415, "Upload your resume as a PDF or a Word (.docx) file.")


def _is_docx(data: bytes) -> bool:
    if not data.startswith(b"PK"):
        return False
    try:
        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            return "word/document.xml" in archive.namelist()
    except zipfile.BadZipFile:
        return False


def docx_text(data: bytes) -> str:
    import docx

    document = docx.Document(io.BytesIO(data))
    lines = [paragraph.text for paragraph in document.paragraphs]
    # Resumes love tables for layout; their text matters as much as the paragraphs'.
    for table in document.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if cells:
                lines.append(" | ".join(dict.fromkeys(cells)))
    return "\n".join(line for line in lines if line.strip())


def response_text(response) -> str:
    """The text of a Messages API response, after checking it actually finished."""
    if getattr(response, "stop_reason", None) == "refusal":
        raise AIFailed("The AI declined this request.")
    text = "".join(
        block.text for block in response.content if getattr(block, "type", None) == "text"
    ).strip()
    if not text:
        raise AIFailed("The AI returned an empty reply. Try again.")
    return text
